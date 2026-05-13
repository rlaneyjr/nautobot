"""
JOBS_ROOT job: Presidio e-mail Excel device import.

Poll IMAP using ``NB_INGEST_*`` worker environment variables, read ``.xlsx`` attachments, parse the worksheet named
for the current calendar year (or an overridden year), apply Presidio column/serial rules, fuzzy-match ``DeviceType``,
and create ``Device`` records via the same CSV serializer path as Import Objects.

**Deploying:** place this module under ``JOBS_ROOT`` (``NAUTOBOT_JOBS_ROOT`` / ``JOBS_ROOT``), restart Celery workers,
run ``nautobot-server post_upgrade``, enable the Job under **Jobs**.

**Scheduling:** **Jobs > Scheduled Jobs** (Celery Beat running).

**Credentials (environment on workers):**

* ``NB_INGEST_EMAIL`` — mailbox address (used as login when ``NB_INGEST_USERNAME`` is unset).
* ``NB_INGEST_USERNAME`` — IMAP login; if unset, ``NB_INGEST_EMAIL`` is used.
* ``NB_INGEST_PASSWORD`` — IMAP password.
* ``NB_INGEST_IMAP_HOST`` — when set, overrides the Job form ``imap_host`` field.
"""

from __future__ import annotations

import codecs
import contextlib
import csv
import difflib
import email
import imaplib
import io
import os
import re
from datetime import datetime
from email.message import Message
from io import BytesIO
from typing import Iterable

import openpyxl
from django.core.exceptions import PermissionDenied
from django.db import transaction

from nautobot.apps.jobs import (
    BooleanVar,
    DryRunVar,
    IntegerVar,
    Job,
    RunJobTaskFailed,
    StringVar,
    TextVar,
    register_jobs,
)
from nautobot.core.api.parsers import NautobotCSVParser
from nautobot.core.api.utils import get_serializer_for_model
from nautobot.core.exceptions import AbortTransaction
from nautobot.dcim.models import Device, DeviceType
from nautobot.extras.models import Role, Status

name = "Presidio email import"

NB_INGEST_EMAIL = "NB_INGEST_EMAIL"
NB_INGEST_USERNAME = "NB_INGEST_USERNAME"
NB_INGEST_PASSWORD = "NB_INGEST_PASSWORD"
NB_INGEST_IMAP_HOST = "NB_INGEST_IMAP_HOST"

PRESIDIO_CARE_COLUMNS = ["Date Placed", "PO #", "Description", "Location", "SKU", "QTY", "Serial Numbers"]
PRESIDIO_REQUIRED_COLUMNS = ["Description", "Location", "SKU", "Serial Numbers"]
PRESIDIO_DEFAULT_STATUS_NAME = "Inventory"
PRESIDIO_DEFAULT_ROLE_NAME = "unknown"
PRESIDIO_SERIAL_CONTRACT_AVOIDANCE = "contract #"
PRESIDIO_DATE_FORMATS = ("%m/%d/%Y", "%m/%d/%y")


def connect_imap_client(*, host, port, use_ssl, username, password):
    """Connect and log in to IMAP; raise ``RunJobTaskFailed`` on failure."""
    if use_ssl:
        client = imaplib.IMAP4_SSL(host, port)
    else:
        client = imaplib.IMAP4(host, port)
    try:
        typ, _ = client.login(username, password)
        if typ != "OK":
            raise RunJobTaskFailed(f"IMAP login failed: {typ}")
    except imaplib.IMAP4.error as exc:
        raise RunJobTaskFailed(f"IMAP login error: {exc}") from exc
    return client


def resolve_presidio_imap_host(*, env_host: str, job_imap_host: str) -> str:
    h = (env_host or "").strip()
    if h:
        return h
    h = (job_imap_host or "").strip()
    if h:
        return h
    raise RunJobTaskFailed(
        f"Set {NB_INGEST_IMAP_HOST} in the worker environment or provide IMAP hostname on the Job form."
    )


def resolve_presidio_imap_credentials():
    """Return (username, password) from NB_INGEST_* environment variables."""
    email_addr = (os.getenv(NB_INGEST_EMAIL) or "").strip()
    username = (os.getenv(NB_INGEST_USERNAME) or "").strip() or email_addr
    password = (os.getenv(NB_INGEST_PASSWORD) or "").strip()
    if not username or not password:
        raise RunJobTaskFailed(
            f"Set {NB_INGEST_USERNAME} (or {NB_INGEST_EMAIL}) and {NB_INGEST_PASSWORD} for this job."
        )
    return username, password


def device_import_perform_operation(logger, user, data, serializer_class, queryset):
    new_objs = []
    validation_failed = False
    for row, entry in enumerate(data, start=1):
        serializer = serializer_class(data=entry, context={"request": None})
        if serializer.is_valid():
            try:
                with transaction.atomic():
                    new_obj = serializer.save()
                    if not queryset.filter(pk=new_obj.pk).exists():
                        raise AbortTransaction()
                logger.info('Row %d: Created record "%s"', row, new_obj, extra={"object": new_obj})
                new_objs.append(new_obj)
            except AbortTransaction:
                logger.error(
                    'Row %d: User "%s" does not have permission to create an object with these attributes',
                    row,
                    user,
                )
                validation_failed = True
        else:
            validation_failed = True
            for field, errs in serializer.errors.items():
                for err in errs:
                    logger.error("Row %d: `%s`: `%s`", row, field, err)
    return new_objs, validation_failed


def device_import_perform_atomic_operation(logger, user, data, serializer_class, queryset):
    new_objs = []
    with contextlib.suppress(AbortTransaction):
        with transaction.atomic():
            new_objs, validation_failed = device_import_perform_operation(
                logger, user, data, serializer_class, queryset
            )
            if validation_failed:
                raise AbortTransaction()
            return new_objs, validation_failed
    logger.warning("Rolling back all %s records from this attachment.", len(new_objs))
    return [], True


def import_devices_from_normalized_csv_bytes(
    logger,
    user,
    csv_bytes: bytes,
    *,
    roll_back_if_error: bool,
    dryrun: bool,
    skip_existing_serial: bool,
):
    if not user.has_perm("dcim.add_device"):
        raise PermissionDenied("User does not have permission to create Device records")

    serializer_class = get_serializer_for_model(Device)
    queryset = Device.objects.restrict(user, "add")
    stream = BytesIO(csv_bytes)
    try:
        data = NautobotCSVParser().parse(
            stream=stream,
            parser_context={"request": None, "serializer_class": serializer_class},
        )
    except Exception as exc:
        logger.error("CSV parse error: %s", exc)
        raise RunJobTaskFailed(f"CSV parse failed: {exc}") from exc

    if skip_existing_serial:
        filtered = []
        for entry in data:
            serial = None
            if isinstance(entry, dict):
                serial = (entry.get("serial") or "").strip()
            if serial and Device.objects.restrict(user, "view").filter(serial=serial).exists():
                logger.info("Skipping row with existing serial %s", serial)
                continue
            filtered.append(entry)
        data = filtered

    if dryrun:
        logger.info("Dry-run: validated parse produced %d row(s); skipping database writes.", len(data))
        return []

    if roll_back_if_error:
        new_objs, validation_failed = device_import_perform_atomic_operation(
            logger, user, data, serializer_class, queryset
        )
    else:
        new_objs, validation_failed = device_import_perform_operation(logger, user, data, serializer_class, queryset)

    if validation_failed:
        if roll_back_if_error:
            raise RunJobTaskFailed("CSV import not successful; rolled back this attachment")
        raise RunJobTaskFailed("CSV import not fully successful for this attachment; see logs")

    return new_objs


def normalized_rows_to_csv_bytes(rows: Iterable[dict[str, str]]) -> bytes:
    rows = list(rows)
    if not rows:
        return codecs.BOM_UTF8
    fieldnames: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r:
            if k not in seen:
                seen.add(k)
                fieldnames.append(k)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow({k: r.get(k, "") for k in fieldnames})
    return codecs.BOM_UTF8 + buf.getvalue().encode("utf-8")


def extract_xlsx_attachments_from_message(msg: Message, *, max_bytes: int) -> list[tuple[str | None, bytes]]:
    xlsx_ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    found: list[tuple[str | None, bytes]] = []
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_maintype() == "multipart":
                continue
            filename = part.get_filename()
            ctype = part.get_content_type()
            is_xlsx = ctype == xlsx_ct or (filename and filename.lower().endswith(".xlsx"))
            if not is_xlsx:
                continue
            raw = part.get_payload(decode=True)
            if raw is None:
                continue
            if len(raw) > max_bytes:
                raise ValueError(f"Attachment exceeds max size ({len(raw)} > {max_bytes})")
            found.append((filename, raw))
    else:
        raw = msg.get_payload(decode=True)
        if raw and len(raw) <= max_bytes:
            fn = msg.get_filename()
            if fn and fn.lower().endswith(".xlsx"):
                found.append((fn, raw))
    return found


def split_presidio_serial_tokens(serial_cell: str) -> list[str]:
    if not serial_cell or not str(serial_cell).strip():
        return []
    parts = re.split(r",\s*|\s+", str(serial_cell).strip())
    out: list[str] = []
    for p in parts:
        t = p.strip()
        if not t:
            continue
        low = t.lower()
        if low == PRESIDIO_SERIAL_CONTRACT_AVOIDANCE or low.startswith(PRESIDIO_SERIAL_CONTRACT_AVOIDANCE):
            continue
        out.append(t)
    return out


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_presidio_qty(qty_val) -> int | None:
    if qty_val is None or qty_val == "":
        return None
    try:
        return int(float(str(qty_val).strip()))
    except (TypeError, ValueError):
        return None


def _row_has_nondefault_fill(ws_style, row_idx: int, col_indices: Iterable[int]) -> bool:
    for col in col_indices:
        cell = ws_style.cell(row=row_idx, column=col)
        fill = cell.fill
        if fill is None:
            continue
        ft = fill.fill_type
        if ft is None:
            continue
        if str(ft).lower() == "none":
            continue
        return True
    return False


def find_presidio_header_row_and_columns(ws, *, max_scan_rows: int = 50) -> tuple[int, dict[str, int]]:
    required_set = set(PRESIDIO_REQUIRED_COLUMNS)
    for row_idx in range(1, max_scan_rows + 1):
        headers: dict[str, int] = {}
        max_col = min(ws.max_column or 1, 512)
        for col_idx in range(1, max_col + 1):
            val = _cell_str(ws.cell(row=row_idx, column=col_idx).value)
            if val in PRESIDIO_CARE_COLUMNS:
                headers[val] = col_idx
        if required_set.issubset(headers.keys()):
            return row_idx, headers
    raise ValueError(
        f"Could not find a header row containing required columns {sorted(required_set)!r} in first {max_scan_rows} rows."
    )


def _read_presidio_row(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, str]:
    return {name: _cell_str(ws.cell(row=row_idx, column=col).value) for name, col in col_map.items()}


def _presidio_row_required_all_present(raw: dict[str, str]) -> bool:
    return all((raw.get(c) or "").strip() for c in PRESIDIO_REQUIRED_COLUMNS)


def _presidio_merge_logical(context: dict[str, str], raw: dict[str, str]) -> dict[str, str]:
    return {c: (raw.get(c) or "").strip() or context.get(c, "") for c in PRESIDIO_CARE_COLUMNS}


def _parse_date_placed_for_log(raw_date: str, logger) -> None:
    if not (raw_date or "").strip():
        return
    for fmt in PRESIDIO_DATE_FORMATS:
        try:
            datetime.strptime(raw_date.strip(), fmt)
            return
        except ValueError:
            continue
    logger.warning("Date Placed %r did not match MM/DD/YYYY (tried %s).", raw_date, PRESIDIO_DATE_FORMATS)


def best_presidio_device_type_match(
    device_types: list[DeviceType],
    *,
    sku: str,
    description: str,
    min_ratio: float,
) -> tuple[DeviceType | None, float]:
    query = f"{sku} {description}".lower()
    query = " ".join(query.split())
    if not query.strip():
        return None, 0.0
    best_dt: DeviceType | None = None
    best_score = 0.0
    for dt in device_types:
        candidates = [dt.model, dt.part_number, f"{dt.model} {dt.part_number}".strip()]
        if getattr(dt, "comments", None):
            candidates.append(dt.comments)
        for candidate in candidates:
            if not candidate:
                continue
            cand = " ".join(str(candidate).lower().split())
            score = difflib.SequenceMatcher(None, query, cand).ratio()
            if score > best_score:
                best_score = score
                best_dt = dt
    if best_dt is not None and best_score >= min_ratio:
        return best_dt, best_score
    return None, best_score


def presidio_xlsx_bytes_to_normalized_device_rows(
    xlsx_bytes: bytes,
    *,
    sheet_name: str,
    logger,
    user,
    status_name: str,
    role_name: str,
    fuzzy_min_score: float,
    fail_on_unmatched_device_type: bool,
) -> list[dict[str, str]]:
    raw_bio = BytesIO(xlsx_bytes)
    wb_values = openpyxl.load_workbook(raw_bio, data_only=True)
    if sheet_name not in wb_values.sheetnames:
        names = list(wb_values.sheetnames)
        wb_values.close()
        raise RunJobTaskFailed(f"Worksheet {sheet_name!r} not found. Available sheets: {names!r}")
    ws_values = wb_values[sheet_name]

    try:
        header_row_idx, col_map = find_presidio_header_row_and_columns(ws_values)
    except ValueError as exc:
        wb_values.close()
        raise RunJobTaskFailed(str(exc)) from exc

    raw_bio2 = BytesIO(xlsx_bytes)
    wb_style = openpyxl.load_workbook(raw_bio2, data_only=False)
    ws_style = wb_style[sheet_name]

    style_col_indices = list(col_map.values())
    device_types = list(
        DeviceType.objects.restrict(user, "view").select_related("manufacturer").iterator(chunk_size=500)
    )

    context: dict[str, str] = {c: "" for c in PRESIDIO_CARE_COLUMNS}
    seen_first_complete = False
    last_emitted_serial: str | None = None
    flat_rows: list[dict[str, str]] = []

    try:
        for row_idx in range(header_row_idx + 1, ws_values.max_row + 1):
            raw = _read_presidio_row(ws_values, row_idx, col_map)
            required_empty = not any((raw.get(c) or "").strip() for c in PRESIDIO_REQUIRED_COLUMNS)
            shaded = _row_has_nondefault_fill(ws_style, row_idx, style_col_indices)
            if required_empty and shaded:
                continue
            if required_empty and not seen_first_complete:
                continue

            if not seen_first_complete:
                if not _presidio_row_required_all_present(raw):
                    continue
                seen_first_complete = True
                context = {c: (raw.get(c) or "").strip() for c in PRESIDIO_CARE_COLUMNS}
                logical = context.copy()
            else:
                logical = _presidio_merge_logical(context, raw)
                if _presidio_row_required_all_present(raw):
                    context = logical.copy()

            _parse_date_placed_for_log(logical.get("Date Placed", ""), logger)

            serials = split_presidio_serial_tokens(logical.get("Serial Numbers", ""))
            qty = _parse_presidio_qty(logical.get("QTY"))
            if qty is not None and qty > len(serials):
                logger.warning(
                    "Row %s: QTY=%s exceeds parsed serial count %s (PO=%r SKU=%r).",
                    row_idx,
                    qty,
                    len(serials),
                    logical.get("PO #", ""),
                    logical.get("SKU", ""),
                )

            for sn in serials:
                if last_emitted_serial is not None and sn == last_emitted_serial:
                    continue
                last_emitted_serial = sn

                if Device.objects.restrict(user, "view").filter(serial=sn).exists():
                    logger.info("Skipping serial %s: already present on a Device", sn)
                    continue

                sku = logical.get("SKU", "")
                description = logical.get("Description", "")
                dt, score = best_presidio_device_type_match(
                    device_types, sku=sku, description=description, min_ratio=fuzzy_min_score
                )
                if dt is None:
                    msg = (
                        f"No DeviceType match for serial {sn!r} (SKU={sku!r}, description={description!r}; "
                        f"best score {score:.3f} below threshold {fuzzy_min_score})."
                    )
                    if fail_on_unmatched_device_type:
                        raise RunJobTaskFailed(msg)
                    logger.error("%s Skipping.", msg)
                    continue

                flat_rows.append(
                    {
                        "serial": sn,
                        "name": sn,
                        "device_type__manufacturer__name": dt.manufacturer.name,
                        "device_type__model": dt.model,
                        "location__name": logical.get("Location", ""),
                        "status__name": status_name,
                        "role__name": role_name,
                    }
                )
        return flat_rows
    finally:
        wb_values.close()
        wb_style.close()


def _msg_set(msg_id: bytes | str) -> str:
    return msg_id.decode("ascii") if isinstance(msg_id, bytes) else str(msg_id)


def _sender_allowed(from_header: str, allowlist: str) -> bool:
    if not allowlist.strip():
        return True
    emails = {e.strip().lower() for e in allowlist.split(",") if e.strip()}
    if not emails:
        return True
    candidates = re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", from_header or "")
    lowered = {c.lower() for c in candidates}
    return bool(lowered & emails)


def resolve_presidio_inventory_status_and_unknown_role_names():
    status = Status.objects.get_for_model(Device).filter(name__iexact=PRESIDIO_DEFAULT_STATUS_NAME).first()
    if not status:
        raise RunJobTaskFailed(
            f'Device Status named "{PRESIDIO_DEFAULT_STATUS_NAME}" is required for the Presidio e-mail import job.'
        )
    role = Role.objects.get_for_model(Device).filter(name__iexact=PRESIDIO_DEFAULT_ROLE_NAME).first()
    if not role:
        raise RunJobTaskFailed(
            f'Device Role named "{PRESIDIO_DEFAULT_ROLE_NAME}" is required for the Presidio e-mail import job.'
        )
    return status.name, role.name


class PresidioEmailImportJob(Job):
    """Presidio: IMAP + Excel workbook → Devices (NB_INGEST_* env credentials)."""

    imap_host = StringVar(
        required=False,
        default="",
        description=f"IMAP hostname when {NB_INGEST_IMAP_HOST} is not set in the worker environment",
    )
    imap_port = IntegerVar(default=993, description="IMAP port (993 for SSL)")
    imap_use_ssl = BooleanVar(default=True, description="Use IMAP4_SSL (disable for cleartext / custom setups)")
    imap_mailbox = StringVar(default="INBOX", description="Mailbox folder name")
    imap_search_criteria = StringVar(
        default="UNSEEN",
        description='IMAP SEARCH criteria (e.g. "UNSEEN" or \'UNSEEN FROM "sender@example.com"\')',
    )
    worksheet_year = IntegerVar(
        default=0,
        description="Worksheet tab name as a year (e.g. 2026). Use 0 for the current calendar year.",
    )
    fuzzy_match_min_percent = IntegerVar(
        default=42,
        description="Minimum fuzzy match score for DeviceType (1-100, default 42 means 0.42).",
    )
    fail_on_unmatched_device_type = BooleanVar(
        default=False,
        description="If true, fail the job when a serial row has no DeviceType above the fuzzy threshold.",
    )
    sender_allowlist = TextVar(
        required=False,
        description="Optional comma-separated sender e-mail addresses to accept (empty = any)",
    )
    max_messages = IntegerVar(default=25, description="Maximum number of messages to process in one run")
    max_attachment_bytes = IntegerVar(
        default=5_000_000,
        description="Maximum Excel attachment size in bytes",
    )
    roll_back_if_error = BooleanVar(
        default=True,
        description="If any row fails validation, roll back the entire import for that Excel attachment",
    )
    skip_existing_serial = BooleanVar(
        default=True,
        description="Skip rows whose serial already exists on a Device",
    )
    mark_seen_on_success = BooleanVar(
        default=True,
        description=r"Set \Seen on the message after the attachment imports successfully",
    )
    move_to_mailbox = StringVar(
        required=False,
        description="If set, COPY the message to this mailbox after success, then delete from the source mailbox",
    )
    dryrun = DryRunVar()

    class Meta:
        name = "Presidio email import"
        description = (
            "Presidio: connect with NB_INGEST_* credentials, find messages with .xlsx attachments, parse the year "
            "worksheet, fuzzy-match DeviceType, and create Devices via the Import Objects CSV validation path."
        )
        has_sensitive_variables = False
        soft_time_limit = 1800
        time_limit = 2000

    def run(  # pylint: disable=too-many-locals,too-many-branches,too-many-statements
        self,
        *,
        imap_host="",
        imap_port=993,
        imap_use_ssl=True,
        imap_mailbox="INBOX",
        imap_search_criteria="UNSEEN",
        worksheet_year=0,
        fuzzy_match_min_percent=42,
        fail_on_unmatched_device_type=False,
        sender_allowlist="",
        max_messages=25,
        max_attachment_bytes=5_000_000,
        roll_back_if_error=True,
        skip_existing_serial=True,
        mark_seen_on_success=True,
        move_to_mailbox="",
        dryrun=False,
    ):
        host = resolve_presidio_imap_host(env_host=os.getenv(NB_INGEST_IMAP_HOST, ""), job_imap_host=imap_host or "")
        imap_username, password = resolve_presidio_imap_credentials()
        client = connect_imap_client(
            host=host,
            port=imap_port,
            use_ssl=imap_use_ssl,
            username=imap_username,
            password=password,
        )
        status_name, role_name = resolve_presidio_inventory_status_and_unknown_role_names()
        year = worksheet_year if worksheet_year else datetime.now().year
        sheet_name = str(year)
        pct = max(1, min(100, int(fuzzy_match_min_percent)))
        fuzzy_min_score = pct / 100.0

        try:
            typ, _ = client.select(imap_mailbox)
            if typ != "OK":
                raise RunJobTaskFailed(f"Unable to select mailbox {imap_mailbox!r}: {typ}")

            typ, data = client.uid("search", None, imap_search_criteria)
            if typ != "OK":
                raise RunJobTaskFailed(f"IMAP UID SEARCH failed: {typ}")

            raw_ids = data[0] if data else None
            if not raw_ids:
                msg_uids = []
            else:
                msg_uids = raw_ids.split()
            if not msg_uids:
                self.logger.info("No messages matched search %s", imap_search_criteria)
                return

            if len(msg_uids) > max_messages:
                self.logger.warning("Truncating to max_messages=%s (had %s)", max_messages, len(msg_uids))
                msg_uids = msg_uids[:max_messages]

            move_mb = str(move_to_mailbox).strip() if move_to_mailbox else ""

            for uid in msg_uids:
                uid_str = _msg_set(uid)
                typ, msg_data = client.uid("fetch", uid_str, "(RFC822)")
                if typ != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
                    self.logger.error("FETCH failed for message UID %s", uid_str)
                    continue
                raw_msg = msg_data[0][1]
                msg = email.message_from_bytes(raw_msg)
                from_header = msg.get("From", "")
                if not _sender_allowed(from_header, sender_allowlist or ""):
                    self.logger.info("Skipping message UID %s: sender not in allowlist", uid_str)
                    continue

                attachments = extract_xlsx_attachments_from_message(msg, max_bytes=max_attachment_bytes)
                if not attachments:
                    self.logger.info("No Excel attachments on message UID %s", uid_str)
                    continue

                message_import_completed = False
                for filename, raw_xlsx in attachments:
                    self.logger.info(
                        "Processing Excel attachment %s from message UID %s (sheet %r)",
                        filename or "(no filename)",
                        uid_str,
                        sheet_name,
                    )
                    flat_rows = presidio_xlsx_bytes_to_normalized_device_rows(
                        raw_xlsx,
                        sheet_name=sheet_name,
                        logger=self.logger,
                        user=self.user,
                        status_name=status_name,
                        role_name=role_name,
                        fuzzy_min_score=fuzzy_min_score,
                        fail_on_unmatched_device_type=fail_on_unmatched_device_type,
                    )
                    if not flat_rows:
                        self.logger.info("Attachment %s on UID %s produced no device rows; skipping", filename, uid_str)
                        continue

                    csv_out = normalized_rows_to_csv_bytes(flat_rows)
                    try:
                        import_devices_from_normalized_csv_bytes(
                            self.logger,
                            self.user,
                            csv_out,
                            roll_back_if_error=roll_back_if_error,
                            dryrun=dryrun,
                            skip_existing_serial=skip_existing_serial,
                        )
                    except RunJobTaskFailed:
                        self.logger.error("Import failed for attachment %s on message UID %s", filename, uid_str)
                        raise

                    message_import_completed = True

                if message_import_completed and mark_seen_on_success and not dryrun:
                    client.uid("store", uid_str, "+FLAGS", r"\Seen")

                if message_import_completed and move_mb and not dryrun:
                    typ, _ = client.uid("copy", uid_str, move_mb)
                    if typ != "OK":
                        raise RunJobTaskFailed(f"IMAP UID COPY to {move_mb!r} failed: {typ}")
                    client.uid("store", uid_str, "+FLAGS", r"\Deleted")
                    self.logger.info("Moved message UID %s to mailbox %s (copy+delete source)", uid_str, move_mb)

            if move_mb and not dryrun:
                client.expunge()

        finally:
            try:
                client.logout()
            except imaplib.IMAP4.error:
                pass


register_jobs(PresidioEmailImportJob)
