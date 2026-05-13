"""
JOBS_ROOT job: RedThread e-mail Excel device import.

Poll IMAP using ``NB_REDTHREAD_*`` worker environment variables, read ``.xlsx`` attachments, parse a worksheet
(tab name = calendar year or override), map straightforward Excel columns to Nautobot Device CSV keys, and create
``Device`` records via the same serializer path as Import Objects.

**Excel columns** (header row; case/spacing-insensitive, including common typo ``Manufaturer``):

* **Manufacturer** / Manufaturer → ``device_type__manufacturer__name``
* **Model** → ``device_type__model`` (if empty, **DeviceType** cell is used as model)
* **SerialNumber** / Serial Number → ``serial`` and ``name``
* **Location** → ``location__name`` (blank defaults to **RedThread Warehouse**, overridable on the Job form)

**Deploying:** place under ``JOBS_ROOT``, restart Celery workers, ``nautobot-server post_upgrade``, enable under **Jobs**.

**Credentials (worker environment):** ``NB_REDTHREAD_EMAIL``, ``NB_REDTHREAD_USERNAME``, ``NB_REDTHREAD_PASSWORD``,
``NB_REDTHREAD_IMAP_HOST`` (host env overrides Job ``imap_host`` when set).
"""

from __future__ import annotations

import codecs
import contextlib
import csv
import email
import imaplib
import io
import os
import re
from datetime import datetime
from email.message import Message
from io import BytesIO
from typing import Iterable

import openpyxl
from django.core.exceptions import PermissionDenied
from django.db import transaction

from nautobot.apps.jobs import (
    BooleanVar,
    DryRunVar,
    IntegerVar,
    Job,
    RunJobTaskFailed,
    StringVar,
    TextVar,
    register_jobs,
)
from nautobot.core.api.parsers import NautobotCSVParser
from nautobot.core.api.utils import get_serializer_for_model
from nautobot.core.exceptions import AbortTransaction
from nautobot.dcim.models import Device, DeviceType
from nautobot.extras.models import Role, Status

name = "RedThread email import"

NB_REDTHREAD_EMAIL = "NB_REDTHREAD_EMAIL"
NB_REDTHREAD_USERNAME = "NB_REDTHREAD_USERNAME"
NB_REDTHREAD_PASSWORD = "NB_REDTHREAD_PASSWORD"
NB_REDTHREAD_IMAP_HOST = "NB_REDTHREAD_IMAP_HOST"

REDTHREAD_DEFAULT_LOCATION_NAME = "RedThread Warehouse"
REDTHREAD_DEFAULT_STATUS_NAME = "Inventory"
REDTHREAD_DEFAULT_ROLE_NAME = "unknown"


def _normalize_excel_header_label(cell_value) -> str | None:
    """Map a raw header cell to a canonical logical name, or None if unknown."""
    if cell_value is None:
        return None
    s = str(cell_value).strip()
    if not s:
        return None
    key = "".join(s.split()).lower()
    if key in ("manufacturer", "manufaturer"):
        return "Manufacturer"
    if key == "model":
        return "Model"
    if key in ("devicetype", "device_type"):
        return "DeviceType"
    if key in ("serialnumber", "serial_number", "serial"):
        return "SerialNumber"
    if key == "location":
        return "Location"
    return None


def connect_imap_client(*, host, port, use_ssl, username, password):
    if use_ssl:
        client = imaplib.IMAP4_SSL(host, port)
    else:
        client = imaplib.IMAP4(host, port)
    try:
        typ, _ = client.login(username, password)
        if typ != "OK":
            raise RunJobTaskFailed(f"IMAP login failed: {typ}")
    except imaplib.IMAP4.error as exc:
        raise RunJobTaskFailed(f"IMAP login error: {exc}") from exc
    return client


def resolve_redthread_imap_host(*, env_host: str, job_imap_host: str) -> str:
    h = (env_host or "").strip()
    if h:
        return h
    h = (job_imap_host or "").strip()
    if h:
        return h
    raise RunJobTaskFailed(
        f"Set {NB_REDTHREAD_IMAP_HOST} in the worker environment or provide IMAP hostname on the Job form."
    )


def resolve_redthread_imap_credentials():
    email_addr = (os.getenv(NB_REDTHREAD_EMAIL) or "").strip()
    username = (os.getenv(NB_REDTHREAD_USERNAME) or "").strip() or email_addr
    password = (os.getenv(NB_REDTHREAD_PASSWORD) or "").strip()
    if not username or not password:
        raise RunJobTaskFailed(
            f"Set {NB_REDTHREAD_USERNAME} (or {NB_REDTHREAD_EMAIL}) and {NB_REDTHREAD_PASSWORD} for this job."
        )
    return username, password


def device_import_perform_operation(logger, user, data, serializer_class, queryset):
    new_objs = []
    validation_failed = False
    for row, entry in enumerate(data, start=1):
        serializer = serializer_class(data=entry, context={"request": None})
        if serializer.is_valid():
            try:
                with transaction.atomic():
                    new_obj = serializer.save()
                    if not queryset.filter(pk=new_obj.pk).exists():
                        raise AbortTransaction()
                logger.info('Row %d: Created record "%s"', row, new_obj, extra={"object": new_obj})
                new_objs.append(new_obj)
            except AbortTransaction:
                logger.error(
                    'Row %d: User "%s" does not have permission to create an object with these attributes',
                    row,
                    user,
                )
                validation_failed = True
        else:
            validation_failed = True
            for field, errs in serializer.errors.items():
                for err in errs:
                    logger.error("Row %d: `%s`: `%s`", row, field, err)
    return new_objs, validation_failed


def device_import_perform_atomic_operation(logger, user, data, serializer_class, queryset):
    new_objs = []
    with contextlib.suppress(AbortTransaction):
        with transaction.atomic():
            new_objs, validation_failed = device_import_perform_operation(
                logger, user, data, serializer_class, queryset
            )
            if validation_failed:
                raise AbortTransaction()
            return new_objs, validation_failed
    logger.warning("Rolling back all %s records from this attachment.", len(new_objs))
    return [], True


def import_devices_from_normalized_csv_bytes(
    logger,
    user,
    csv_bytes: bytes,
    *,
    roll_back_if_error: bool,
    dryrun: bool,
    skip_existing_serial: bool,
):
    if not user.has_perm("dcim.add_device"):
        raise PermissionDenied("User does not have permission to create Device records")

    serializer_class = get_serializer_for_model(Device)
    queryset = Device.objects.restrict(user, "add")
    stream = BytesIO(csv_bytes)
    try:
        data = NautobotCSVParser().parse(
            stream=stream,
            parser_context={"request": None, "serializer_class": serializer_class},
        )
    except Exception as exc:
        logger.error("CSV parse error: %s", exc)
        raise RunJobTaskFailed(f"CSV parse failed: {exc}") from exc

    if skip_existing_serial:
        filtered = []
        for entry in data:
            serial = None
            if isinstance(entry, dict):
                serial = (entry.get("serial") or "").strip()
            if serial and Device.objects.restrict(user, "view").filter(serial=serial).exists():
                logger.info("Skipping row with existing serial %s", serial)
                continue
            filtered.append(entry)
        data = filtered

    if dryrun:
        logger.info("Dry-run: validated parse produced %d row(s); skipping database writes.", len(data))
        return []

    if roll_back_if_error:
        new_objs, validation_failed = device_import_perform_atomic_operation(
            logger, user, data, serializer_class, queryset
        )
    else:
        new_objs, validation_failed = device_import_perform_operation(logger, user, data, serializer_class, queryset)

    if validation_failed:
        if roll_back_if_error:
            raise RunJobTaskFailed("CSV import not successful; rolled back this attachment")
        raise RunJobTaskFailed("CSV import not fully successful for this attachment; see logs")

    return new_objs


def normalized_rows_to_csv_bytes(rows: Iterable[dict[str, str]]) -> bytes:
    rows = list(rows)
    if not rows:
        return codecs.BOM_UTF8
    fieldnames: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r:
            if k not in seen:
                seen.add(k)
                fieldnames.append(k)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        writer.writerow({k: r.get(k, "") for k in fieldnames})
    return codecs.BOM_UTF8 + buf.getvalue().encode("utf-8")


def extract_xlsx_attachments_from_message(msg: Message, *, max_bytes: int) -> list[tuple[str | None, bytes]]:
    xlsx_ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    found: list[tuple[str | None, bytes]] = []
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_maintype() == "multipart":
                continue
            filename = part.get_filename()
            ctype = part.get_content_type()
            is_xlsx = ctype == xlsx_ct or (filename and filename.lower().endswith(".xlsx"))
            if not is_xlsx:
                continue
            raw = part.get_payload(decode=True)
            if raw is None:
                continue
            if len(raw) > max_bytes:
                raise ValueError(f"Attachment exceeds max size ({len(raw)} > {max_bytes})")
            found.append((filename, raw))
    else:
        raw = msg.get_payload(decode=True)
        if raw and len(raw) <= max_bytes:
            fn = msg.get_filename()
            if fn and fn.lower().endswith(".xlsx"):
                found.append((fn, raw))
    return found


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def find_redthread_header_row_and_columns(ws, *, max_scan_rows: int = 50) -> tuple[int, dict[str, int]]:
    """
    Find a row where normalized headers include Manufacturer, SerialNumber, and at least one of Model / DeviceType.
    Returns (1-based header row index, map canonical_name -> 1-based column index).
    """
    need = {"Manufacturer", "SerialNumber"}
    model_source = frozenset({"Model", "DeviceType"})
    for row_idx in range(1, max_scan_rows + 1):
        col_map: dict[str, int] = {}
        max_col = min(ws.max_column or 1, 512)
        for col_idx in range(1, max_col + 1):
            canon = _normalize_excel_header_label(ws.cell(row=row_idx, column=col_idx).value)
            if canon and canon not in col_map:
                col_map[canon] = col_idx
        if not need.issubset(col_map.keys()):
            continue
        if not (model_source & col_map.keys()):
            continue
        return row_idx, col_map
    raise ValueError(
        "Could not find a header row with Manufacturer, SerialNumber, and Model or DeviceType "
        f"(scanned first {max_scan_rows} rows)."
    )


def _read_redthread_row(ws, row_idx: int, col_map: dict[str, int]) -> dict[str, str]:
    out: dict[str, str] = {}
    for canon, col in col_map.items():
        out[canon] = _cell_str(ws.cell(row=row_idx, column=col).value)
    return out


def resolve_model_and_log_device_type_mismatch(raw: dict[str, str], logger, row_idx: int) -> str | None:
    """Return device_type__model string, or None if missing."""
    model_cell = (raw.get("Model") or "").strip()
    dt_cell = (raw.get("DeviceType") or "").strip()
    if model_cell:
        if dt_cell and dt_cell != model_cell:
            logger.info(
                "Row %s: Model %r and DeviceType %r both set; using Model for Nautobot device_type.",
                row_idx,
                model_cell,
                dt_cell,
            )
        return model_cell
    if dt_cell:
        return dt_cell
    return None


def redthread_xlsx_bytes_to_normalized_device_rows(
    xlsx_bytes: bytes,
    *,
    sheet_name: str,
    logger,
    user,
    status_name: str,
    role_name: str,
    default_location_name: str,
    fail_on_missing_device_type: bool,
) -> list[dict[str, str]]:
    raw_bio = BytesIO(xlsx_bytes)
    wb_values = openpyxl.load_workbook(raw_bio, data_only=True)
    if sheet_name not in wb_values.sheetnames:
        names = list(wb_values.sheetnames)
        wb_values.close()
        raise RunJobTaskFailed(f"Worksheet {sheet_name!r} not found. Available sheets: {names!r}")
    ws_values = wb_values[sheet_name]

    try:
        header_row_idx, col_map = find_redthread_header_row_and_columns(ws_values)
    except ValueError as exc:
        wb_values.close()
        raise RunJobTaskFailed(str(exc)) from exc

    flat_rows: list[dict[str, str]] = []
    try:
        for row_idx in range(header_row_idx + 1, ws_values.max_row + 1):
            raw = _read_redthread_row(ws_values, row_idx, col_map)
            serial = (raw.get("SerialNumber") or "").strip()
            if not serial:
                continue

            manufacturer = (raw.get("Manufacturer") or "").strip()
            model = resolve_model_and_log_device_type_mismatch(raw, logger, row_idx)
            if not manufacturer or not model:
                logger.warning(
                    "Row %s: skipping row missing Manufacturer or Model/DeviceType (serial=%r).",
                    row_idx,
                    serial,
                )
                continue

            if Device.objects.restrict(user, "view").filter(serial=serial).exists():
                logger.info("Skipping serial %s: already present on a Device", serial)
                continue

            dt = (
                DeviceType.objects.restrict(user, "view")
                .select_related("manufacturer")
                .filter(manufacturer__name__iexact=manufacturer, model__iexact=model)
                .first()
            )
            if dt is None:
                msg = (
                    f"No DeviceType for manufacturer={manufacturer!r} model={model!r} "
                    f"(serial={serial!r}, row {row_idx})."
                )
                if fail_on_missing_device_type:
                    raise RunJobTaskFailed(msg)
                logger.error("%s Skipping.", msg)
                continue

            loc = (raw.get("Location") or "").strip() or default_location_name

            flat_rows.append(
                {
                    "serial": serial,
                    "name": serial,
                    "device_type__manufacturer__name": dt.manufacturer.name,
                    "device_type__model": dt.model,
                    "location__name": loc,
                    "status__name": status_name,
                    "role__name": role_name,
                }
            )
        return flat_rows
    finally:
        wb_values.close()


def _msg_set(msg_id: bytes | str) -> str:
    return msg_id.decode("ascii") if isinstance(msg_id, bytes) else str(msg_id)


def _sender_allowed(from_header: str, allowlist: str) -> bool:
    if not allowlist.strip():
        return True
    emails = {e.strip().lower() for e in allowlist.split(",") if e.strip()}
    if not emails:
        return True
    candidates = re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", from_header or "")
    lowered = {c.lower() for c in candidates}
    return bool(lowered & emails)


def resolve_redthread_inventory_status_and_unknown_role_names():
    status = Status.objects.get_for_model(Device).filter(name__iexact=REDTHREAD_DEFAULT_STATUS_NAME).first()
    if not status:
        raise RunJobTaskFailed(
            f'Device Status named "{REDTHREAD_DEFAULT_STATUS_NAME}" is required for the RedThread e-mail import job.'
        )
    role = Role.objects.get_for_model(Device).filter(name__iexact=REDTHREAD_DEFAULT_ROLE_NAME).first()
    if not role:
        raise RunJobTaskFailed(
            f'Device Role named "{REDTHREAD_DEFAULT_ROLE_NAME}" is required for the RedThread e-mail import job.'
        )
    return status.name, role.name


class RedThreadEmailImportJob(Job):
    """RedThread: IMAP + Excel workbook → Devices (NB_REDTHREAD_* env credentials)."""

    imap_host = StringVar(
        required=False,
        default="",
        description=f"IMAP hostname when {NB_REDTHREAD_IMAP_HOST} is not set in the worker environment",
    )
    imap_port = IntegerVar(default=993, description="IMAP port (993 for SSL)")
    imap_use_ssl = BooleanVar(default=True, description="Use IMAP4_SSL (disable for cleartext / custom setups)")
    imap_mailbox = StringVar(default="INBOX", description="Mailbox folder name")
    imap_search_criteria = StringVar(
        default="UNSEEN",
        description='IMAP SEARCH criteria (e.g. "UNSEEN" or \'UNSEEN FROM "sender@example.com"\')',
    )
    worksheet_year = IntegerVar(
        default=0,
        description="Worksheet tab name as a year (e.g. 2026). Use 0 for the current calendar year.",
    )
    default_location_name = StringVar(
        default=REDTHREAD_DEFAULT_LOCATION_NAME,
        description="Location name used when the Excel Location cell is blank",
    )
    fail_on_missing_device_type = BooleanVar(
        default=False,
        description="If true, fail the job when a row has no exact DeviceType match for manufacturer + model",
    )
    sender_allowlist = TextVar(
        required=False,
        description="Optional comma-separated sender e-mail addresses to accept (empty = any)",
    )
    max_messages = IntegerVar(default=25, description="Maximum number of messages to process in one run")
    max_attachment_bytes = IntegerVar(
        default=5_000_000,
        description="Maximum Excel attachment size in bytes",
    )
    roll_back_if_error = BooleanVar(
        default=True,
        description="If any row fails validation, roll back the entire import for that Excel attachment",
    )
    skip_existing_serial = BooleanVar(
        default=True,
        description="Skip rows whose serial already exists on a Device",
    )
    mark_seen_on_success = BooleanVar(
        default=True,
        description=r"Set \Seen on the message after the attachment imports successfully",
    )
    move_to_mailbox = StringVar(
        required=False,
        description="If set, COPY the message to this mailbox after success, then delete from the source mailbox",
    )
    dryrun = DryRunVar()

    class Meta:
        name = "RedThread email import"
        description = (
            "RedThread: connect with NB_REDTHREAD_* credentials, find .xlsx attachments, parse the year worksheet, "
            "map Manufacturer/Model/SerialNumber/Location to Devices, default blank Location to RedThread Warehouse, "
            "exact DeviceType match, Import Objects CSV path."
        )
        has_sensitive_variables = False
        soft_time_limit = 1800
        time_limit = 2000

    def run(  # pylint: disable=too-many-locals,too-many-branches,too-many-statements
        self,
        *,
        imap_host="",
        imap_port=993,
        imap_use_ssl=True,
        imap_mailbox="INBOX",
        imap_search_criteria="UNSEEN",
        worksheet_year=0,
        default_location_name=REDTHREAD_DEFAULT_LOCATION_NAME,
        fail_on_missing_device_type=False,
        sender_allowlist="",
        max_messages=25,
        max_attachment_bytes=5_000_000,
        roll_back_if_error=True,
        skip_existing_serial=True,
        mark_seen_on_success=True,
        move_to_mailbox="",
        dryrun=False,
    ):
        host = resolve_redthread_imap_host(env_host=os.getenv(NB_REDTHREAD_IMAP_HOST, ""), job_imap_host=imap_host or "")
        imap_username, password = resolve_redthread_imap_credentials()
        client = connect_imap_client(
            host=host,
            port=imap_port,
            use_ssl=imap_use_ssl,
            username=imap_username,
            password=password,
        )
        status_name, role_name = resolve_redthread_inventory_status_and_unknown_role_names()
        year = worksheet_year if worksheet_year else datetime.now().year
        sheet_name = str(year)
        loc_default = (default_location_name or "").strip() or REDTHREAD_DEFAULT_LOCATION_NAME

        try:
            typ, _ = client.select(imap_mailbox)
            if typ != "OK":
                raise RunJobTaskFailed(f"Unable to select mailbox {imap_mailbox!r}: {typ}")

            typ, data = client.uid("search", None, imap_search_criteria)
            if typ != "OK":
                raise RunJobTaskFailed(f"IMAP UID SEARCH failed: {typ}")

            raw_ids = data[0] if data else None
            if not raw_ids:
                msg_uids = []
            else:
                msg_uids = raw_ids.split()
            if not msg_uids:
                self.logger.info("No messages matched search %s", imap_search_criteria)
                return

            if len(msg_uids) > max_messages:
                self.logger.warning("Truncating to max_messages=%s (had %s)", max_messages, len(msg_uids))
                msg_uids = msg_uids[:max_messages]

            move_mb = str(move_to_mailbox).strip() if move_to_mailbox else ""

            for uid in msg_uids:
                uid_str = _msg_set(uid)
                typ, msg_data = client.uid("fetch", uid_str, "(RFC822)")
                if typ != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
                    self.logger.error("FETCH failed for message UID %s", uid_str)
                    continue
                raw_msg = msg_data[0][1]
                msg = email.message_from_bytes(raw_msg)
                from_header = msg.get("From", "")
                if not _sender_allowed(from_header, sender_allowlist or ""):
                    self.logger.info("Skipping message UID %s: sender not in allowlist", uid_str)
                    continue

                attachments = extract_xlsx_attachments_from_message(msg, max_bytes=max_attachment_bytes)
                if not attachments:
                    self.logger.info("No Excel attachments on message UID %s", uid_str)
                    continue

                message_import_completed = False
                for filename, raw_xlsx in attachments:
                    self.logger.info(
                        "Processing Excel attachment %s from message UID %s (sheet %r)",
                        filename or "(no filename)",
                        uid_str,
                        sheet_name,
                    )
                    flat_rows = redthread_xlsx_bytes_to_normalized_device_rows(
                        raw_xlsx,
                        sheet_name=sheet_name,
                        logger=self.logger,
                        user=self.user,
                        status_name=status_name,
                        role_name=role_name,
                        default_location_name=loc_default,
                        fail_on_missing_device_type=fail_on_missing_device_type,
                    )
                    if not flat_rows:
                        self.logger.info("Attachment %s on UID %s produced no device rows; skipping", filename, uid_str)
                        continue

                    csv_out = normalized_rows_to_csv_bytes(flat_rows)
                    try:
                        import_devices_from_normalized_csv_bytes(
                            self.logger,
                            self.user,
                            csv_out,
                            roll_back_if_error=roll_back_if_error,
                            dryrun=dryrun,
                            skip_existing_serial=skip_existing_serial,
                        )
                    except RunJobTaskFailed:
                        self.logger.error("Import failed for attachment %s on message UID %s", filename, uid_str)
                        raise

                    message_import_completed = True

                if message_import_completed and mark_seen_on_success and not dryrun:
                    client.uid("store", uid_str, "+FLAGS", r"\Seen")

                if message_import_completed and move_mb and not dryrun:
                    typ, _ = client.uid("copy", uid_str, move_mb)
                    if typ != "OK":
                        raise RunJobTaskFailed(f"IMAP UID COPY to {move_mb!r} failed: {typ}")
                    client.uid("store", uid_str, "+FLAGS", r"\Deleted")
                    self.logger.info("Moved message UID %s to mailbox %s (copy+delete source)", uid_str, move_mb)

            if move_mb and not dryrun:
                client.expunge()

        finally:
            try:
                client.logout()
            except imaplib.IMAP4.error:
                pass


register_jobs(RedThreadEmailImportJob)
