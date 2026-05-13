"""Tests for Presidio e-mail import job (``presidio_email_import``)."""

import logging
import os
from datetime import datetime
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.test import SimpleTestCase

import openpyxl
from openpyxl.styles import PatternFill

from nautobot.apps.testing import TransactionTestCase
from nautobot.core.choices import ColorChoices
from nautobot.dcim.factory import DeviceFactory
from nautobot.dcim.models import Device
from nautobot.extras.jobs import RunJobTaskFailed, get_job
from nautobot.extras.models import Role, Status
from nautobot.extras.test_jobs import presidio_email_import as presidio_module
from nautobot.extras.test_jobs.presidio_email_import import (
    NB_INGEST_IMAP_HOST,
    NB_INGEST_PASSWORD,
    NB_INGEST_USERNAME,
    PRESIDIO_DEFAULT_ROLE_NAME,
    PRESIDIO_DEFAULT_STATUS_NAME,
    PresidioEmailImportJob,
    best_presidio_device_type_match,
    presidio_xlsx_bytes_to_normalized_device_rows,
    resolve_presidio_imap_host,
    resolve_presidio_imap_credentials,
    resolve_presidio_inventory_status_and_unknown_role_names,
    split_presidio_serial_tokens,
)

User = get_user_model()


def _ensure_presidio_status_and_role():
    """Create Inventory / unknown status and role for Device if absent."""
    ct = ContentType.objects.get_for_model(Device)
    st = Status.objects.filter(name__iexact=PRESIDIO_DEFAULT_STATUS_NAME).first()
    if not st:
        st = Status.objects.create(name=PRESIDIO_DEFAULT_STATUS_NAME, color=ColorChoices.COLOR_GREY)
    if not st.content_types.filter(pk=ct.pk).exists():
        st.content_types.add(ct)
    role = Role.objects.filter(name__iexact=PRESIDIO_DEFAULT_ROLE_NAME).first()
    if not role:
        role = Role.objects.create(name=PRESIDIO_DEFAULT_ROLE_NAME, color=ColorChoices.COLOR_GREY)
    if not role.content_types.filter(pk=ct.pk).exists():
        role.content_types.add(ct)


class PresidioSerialSplitTest(SimpleTestCase):
    def test_split_commas_and_spaces(self):
        self.assertEqual(split_presidio_serial_tokens("A, B C"), ["A", "B", "C"])

    def test_contract_avoidance(self):
        self.assertEqual(split_presidio_serial_tokens("X123, Contract #, Y456"), ["X123", "Y456"])
        self.assertEqual(split_presidio_serial_tokens("Contract #999"), [])


class PresidioFuzzyMatchTest(SimpleTestCase):
    def test_prefers_matching_model(self):
        class M:
            name = "Acme"

        class DT:
            manufacturer = M()
            model = "Widget-9000"
            part_number = "W-9K"
            comments = ""

        dt, score = best_presidio_device_type_match(
            [DT()], sku="W-9K", description="Widget-9000 switch", min_ratio=0.25
        )
        self.assertIsNotNone(dt)
        self.assertGreaterEqual(score, 0.25)


class PresidioResolveHelpersTest(SimpleTestCase):
    def test_imap_host_env_over_job(self):
        with mock.patch.dict(os.environ, {NB_INGEST_IMAP_HOST: "imap.env.test"}, clear=False):
            self.assertEqual(
                resolve_presidio_imap_host(env_host=os.environ.get(NB_INGEST_IMAP_HOST, ""), job_imap_host="ignored"),
                "imap.env.test",
            )

    def test_imap_host_job_fallback(self):
        self.assertEqual(
            resolve_presidio_imap_host(env_host="", job_imap_host="imap.form.test"),
            "imap.form.test",
        )

    def test_imap_host_missing_raises(self):
        with self.assertRaises(RunJobTaskFailed):
            resolve_presidio_imap_host(env_host="", job_imap_host="")


class PresidioWorkbookParseTest(TransactionTestCase):
    databases = ("default", "job_logs")

    def setUp(self):
        super().setUp()
        _ensure_presidio_status_and_role()
        self.template = DeviceFactory()
        self.user = User.objects.create_user(
            username="presidio_excel_job_tester",
            password="password",
            is_superuser=True,
        )
        self.logger = logging.getLogger("test_presidio_excel_import")

    def _build_workbook_bytes(self):
        year_sheet = str(datetime.now().year)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = year_sheet
        headers = ["Date Placed", "PO #", "Description", "Location", "SKU", "QTY", "Serial Numbers"]
        for col, h in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=h)
        ws.cell(row=2, column=1, value="01/15/2026")
        ws.cell(row=2, column=2, value="PO-1")
        ws.cell(row=2, column=3, value=self.template.device_type.model)
        ws.cell(row=2, column=4, value=self.template.location.name)
        ws.cell(row=2, column=5, value=self.template.device_type.part_number or "SKU-1")
        ws.cell(row=2, column=6, value=2)
        ws.cell(row=2, column=7, value="SN-ALPHA, SN-BETA")
        gray = PatternFill(fill_type="solid", fgColor="D9D9D9")
        for c in range(1, 8):
            ws.cell(row=3, column=c).fill = gray
        ws.cell(row=4, column=7, value="SN-GAMMA")
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def test_parse_workbook_emits_rows(self):
        raw = self._build_workbook_bytes()
        status_name, role_name = resolve_presidio_inventory_status_and_unknown_role_names()
        rows = presidio_xlsx_bytes_to_normalized_device_rows(
            raw,
            sheet_name=str(datetime.now().year),
            logger=self.logger,
            user=self.user,
            status_name=status_name,
            role_name=role_name,
            fuzzy_min_score=0.25,
            fail_on_unmatched_device_type=False,
        )
        serials = {r["serial"] for r in rows}
        self.assertIn("SN-ALPHA", serials)
        self.assertIn("SN-BETA", serials)
        self.assertIn("SN-GAMMA", serials)


class PresidioEmailImportJobRunTest(TransactionTestCase):
    databases = ("default", "job_logs")

    def setUp(self):
        super().setUp()
        _ensure_presidio_status_and_role()
        self.template = DeviceFactory()
        self.user = User.objects.create_user(
            username="presidio_imap_tester",
            password="password",
            is_superuser=True,
        )

    def _xlsx_mime_bytes(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(datetime.now().year)
        headers = ["Date Placed", "PO #", "Description", "Location", "SKU", "QTY", "Serial Numbers"]
        for col, h in enumerate(headers, start=1):
            ws.cell(row=1, column=col, value=h)
        ws.cell(row=2, column=3, value=self.template.device_type.model)
        ws.cell(row=2, column=4, value=self.template.location.name)
        ws.cell(row=2, column=5, value=self.template.device_type.part_number or "PN-1")
        ws.cell(row=2, column=6, value=1)
        ws.cell(row=2, column=7, value="IMAP-EXCEL-SN-1")
        bio = BytesIO()
        wb.save(bio)
        root_bytes = bio.getvalue()

        from email.mime.application import MIMEApplication
        from email.mime.multipart import MIMEMultipart

        root = MIMEMultipart()
        root["From"] = "vendor@example.com"
        part = MIMEApplication(root_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        part.add_header("Content-Disposition", "attachment", filename="devices.xlsx")
        root.attach(part)
        return root.as_bytes()

    def _mock_imap_client(self, mime_bytes: bytes):
        client = mock.MagicMock()

        def uid_side_effect(command, *args):
            if command.upper() == "SEARCH":
                return ("OK", [b"99"])
            if command.upper() == "FETCH":
                return ("OK", [(None, mime_bytes)])
            if command.upper() == "STORE":
                return ("OK", [b"99"])
            if command.upper() == "COPY":
                return ("OK", [b""])
            return ("BAD", [b"unknown"])

        client.select.return_value = ("OK", [b"1"])
        client.uid.side_effect = uid_side_effect
        client.expunge.return_value = ("OK", [b""])
        client.logout.return_value = ("BYE", [b""])
        return client

    def test_job_registered(self):
        from nautobot.extras.jobs import get_jobs

        get_jobs(reload=True)
        cls = get_job("presidio_email_import.PresidioEmailImportJob")
        self.assertIs(cls, PresidioEmailImportJob)

    @mock.patch.dict(
        os.environ,
        {
            NB_INGEST_IMAP_HOST: "imap.presidio.test",
            NB_INGEST_USERNAME: "ingest@example.com",
            NB_INGEST_PASSWORD: "secret",
        },
        clear=False,
    )
    def test_resolve_credentials_from_env(self):
        u, p = resolve_presidio_imap_credentials()
        self.assertEqual(u, "ingest@example.com")
        self.assertEqual(p, "secret")

    @mock.patch.dict(
        os.environ,
        {
            NB_INGEST_IMAP_HOST: "imap.presidio.test",
            NB_INGEST_USERNAME: "ingest@example.com",
            NB_INGEST_PASSWORD: "secret",
        },
        clear=False,
    )
    @mock.patch.object(presidio_module, "connect_imap_client")
    def test_import_creates_device_from_xlsx(self, mock_connect):
        mime = self._xlsx_mime_bytes()
        mock_connect.return_value = self._mock_imap_client(mime)

        job = PresidioEmailImportJob()
        job.user = self.user
        job.logger = logging.getLogger("test_presidio_excel_run")
        job.run(
            imap_host="",
            imap_port=993,
            imap_use_ssl=True,
            imap_mailbox="INBOX",
            imap_search_criteria="UNSEEN",
            worksheet_year=0,
            fuzzy_match_min_percent=25,
            fail_on_unmatched_device_type=False,
            sender_allowlist="",
            max_messages=25,
            max_attachment_bytes=5_000_000,
            roll_back_if_error=True,
            skip_existing_serial=True,
            mark_seen_on_success=True,
            move_to_mailbox="",
            dryrun=False,
        )

        created = Device.objects.get(serial="IMAP-EXCEL-SN-1")
        self.assertEqual(created.location_id, self.template.location_id)
        self.assertEqual(created.device_type_id, self.template.device_type_id)

    def test_missing_password_raises(self):
        with mock.patch("nautobot.extras.test_jobs.presidio_email_import.os.getenv", return_value=""):
            with self.assertRaises(RunJobTaskFailed):
                resolve_presidio_imap_credentials()
