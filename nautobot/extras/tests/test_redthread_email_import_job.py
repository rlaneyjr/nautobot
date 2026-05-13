"""Tests for RedThread e-mail import job (``redthread_email_import``)."""

import logging
import os
from datetime import datetime
from io import BytesIO
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.test import SimpleTestCase

import openpyxl

from nautobot.apps.testing import TransactionTestCase
from nautobot.core.choices import ColorChoices
from nautobot.dcim.factory import DeviceFactory
from nautobot.dcim.models import Device
from nautobot.extras.jobs import RunJobTaskFailed, get_job
from nautobot.extras.models import Role, Status
from nautobot.extras.test_jobs import redthread_email_import as redthread_module
from nautobot.extras.test_jobs.redthread_email_import import (
    NB_REDTHREAD_IMAP_HOST,
    NB_REDTHREAD_PASSWORD,
    NB_REDTHREAD_USERNAME,
    REDTHREAD_DEFAULT_LOCATION_NAME,
    REDTHREAD_DEFAULT_ROLE_NAME,
    REDTHREAD_DEFAULT_STATUS_NAME,
    RedThreadEmailImportJob,
    find_redthread_header_row_and_columns,
    redthread_xlsx_bytes_to_normalized_device_rows,
    resolve_model_and_log_device_type_mismatch,
    resolve_redthread_imap_host,
    resolve_redthread_imap_credentials,
    resolve_redthread_inventory_status_and_unknown_role_names,
)

User = get_user_model()


def _ensure_redthread_status_and_role():
    ct = ContentType.objects.get_for_model(Device)
    st = Status.objects.filter(name__iexact=REDTHREAD_DEFAULT_STATUS_NAME).first()
    if not st:
        st = Status.objects.create(name=REDTHREAD_DEFAULT_STATUS_NAME, color=ColorChoices.COLOR_GREY)
    if not st.content_types.filter(pk=ct.pk).exists():
        st.content_types.add(ct)
    role = Role.objects.filter(name__iexact=REDTHREAD_DEFAULT_ROLE_NAME).first()
    if not role:
        role = Role.objects.create(name=REDTHREAD_DEFAULT_ROLE_NAME, color=ColorChoices.COLOR_GREY)
    if not role.content_types.filter(pk=ct.pk).exists():
        role.content_types.add(ct)


class RedThreadHeaderNormalizeTest(SimpleTestCase):
    def test_manufaturer_typo_header_detected(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="Manufaturer")
        ws.cell(row=1, column=2, value="Model")
        ws.cell(row=1, column=3, value="Serial Number")
        row_idx, col_map = find_redthread_header_row_and_columns(ws, max_scan_rows=5)
        self.assertEqual(row_idx, 1)
        self.assertIn("Manufacturer", col_map)
        self.assertIn("Model", col_map)
        self.assertIn("SerialNumber", col_map)


class RedThreadModelColumnTest(SimpleTestCase):
    def test_prefers_model_over_device_type(self):
        log = mock.MagicMock()
        out = resolve_model_and_log_device_type_mismatch(
            {"Model": "ABC", "DeviceType": "XYZ"},
            log,
            2,
        )
        self.assertEqual(out, "ABC")
        log.info.assert_called()

    def test_device_type_fallback(self):
        log = mock.MagicMock()
        out = resolve_model_and_log_device_type_mismatch({"Model": "", "DeviceType": "DT-only"}, log, 3)
        self.assertEqual(out, "DT-only")


class RedThreadResolveHelpersTest(SimpleTestCase):
    def test_imap_host_env_over_job(self):
        with mock.patch.dict(os.environ, {NB_REDTHREAD_IMAP_HOST: "imap.rt.test"}, clear=False):
            self.assertEqual(
                resolve_redthread_imap_host(
                    env_host=os.environ.get(NB_REDTHREAD_IMAP_HOST, ""),
                    job_imap_host="ignored",
                ),
                "imap.rt.test",
            )

    def test_imap_host_missing_raises(self):
        with self.assertRaises(RunJobTaskFailed):
            resolve_redthread_imap_host(env_host="", job_imap_host="")


class RedThreadWorkbookParseTest(TransactionTestCase):
    databases = ("default", "job_logs")

    def setUp(self):
        super().setUp()
        _ensure_redthread_status_and_role()
        self.template = DeviceFactory()
        self.user = User.objects.create_user(
            username="redthread_excel_job_tester",
            password="password",
            is_superuser=True,
        )
        self.logger = logging.getLogger("test_redthread_excel_import")

    def _build_workbook_bytes(self, *, blank_location: bool = True):
        year_sheet = str(datetime.now().year)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = year_sheet
        ws.cell(row=1, column=1, value="Manufacturer")
        ws.cell(row=1, column=2, value="Model")
        ws.cell(row=1, column=3, value="SerialNumber")
        ws.cell(row=1, column=4, value="Location")
        ws.cell(row=2, column=1, value=self.template.device_type.manufacturer.name)
        ws.cell(row=2, column=2, value=self.template.device_type.model)
        ws.cell(row=2, column=3, value="RT-SERIAL-001")
        ws.cell(row=2, column=4, value="" if blank_location else "Custom Site")
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue(), blank_location

    def test_blank_location_uses_warehouse_default(self):
        raw, _ = self._build_workbook_bytes(blank_location=True)
        status_name, role_name = resolve_redthread_inventory_status_and_unknown_role_names()
        rows = redthread_xlsx_bytes_to_normalized_device_rows(
            raw,
            sheet_name=str(datetime.now().year),
            logger=self.logger,
            user=self.user,
            status_name=status_name,
            role_name=role_name,
            default_location_name=REDTHREAD_DEFAULT_LOCATION_NAME,
            fail_on_missing_device_type=False,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["location__name"], REDTHREAD_DEFAULT_LOCATION_NAME)
        self.assertEqual(rows[0]["serial"], "RT-SERIAL-001")

    def test_explicit_location_preserved(self):
        raw, _ = self._build_workbook_bytes(blank_location=False)
        status_name, role_name = resolve_redthread_inventory_status_and_unknown_role_names()
        rows = redthread_xlsx_bytes_to_normalized_device_rows(
            raw,
            sheet_name=str(datetime.now().year),
            logger=self.logger,
            user=self.user,
            status_name=status_name,
            role_name=role_name,
            default_location_name=REDTHREAD_DEFAULT_LOCATION_NAME,
            fail_on_missing_device_type=False,
        )
        self.assertEqual(rows[0]["location__name"], "Custom Site")


class RedThreadEmailImportJobRunTest(TransactionTestCase):
    databases = ("default", "job_logs")

    def setUp(self):
        super().setUp()
        _ensure_redthread_status_and_role()
        self.template = DeviceFactory()
        self.user = User.objects.create_user(
            username="redthread_imap_tester",
            password="password",
            is_superuser=True,
        )

    def _xlsx_mime_bytes(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(datetime.now().year)
        ws.cell(row=1, column=1, value="Manufacturer")
        ws.cell(row=1, column=2, value="Model")
        ws.cell(row=1, column=3, value="SerialNumber")
        ws.cell(row=1, column=4, value="Location")
        ws.cell(row=2, column=1, value=self.template.device_type.manufacturer.name)
        ws.cell(row=2, column=2, value=self.template.device_type.model)
        ws.cell(row=2, column=3, value="RT-IMAP-SN-1")
        ws.cell(row=2, column=4, value=self.template.location.name)
        bio = BytesIO()
        wb.save(bio)
        root_bytes = bio.getvalue()

        from email.mime.application import MIMEApplication
        from email.mime.multipart import MIMEMultipart

        root = MIMEMultipart()
        root["From"] = "vendor@example.com"
        part = MIMEApplication(root_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        part.add_header("Content-Disposition", "attachment", filename="devices.xlsx")
        root.attach(part)
        return root.as_bytes()

    def _mock_imap_client(self, mime_bytes: bytes):
        client = mock.MagicMock()

        def uid_side_effect(command, *args):
            if command.upper() == "SEARCH":
                return ("OK", [b"101"])
            if command.upper() == "FETCH":
                return ("OK", [(None, mime_bytes)])
            if command.upper() == "STORE":
                return ("OK", [b"101"])
            if command.upper() == "COPY":
                return ("OK", [b""])
            return ("BAD", [b"unknown"])

        client.select.return_value = ("OK", [b"1"])
        client.uid.side_effect = uid_side_effect
        client.expunge.return_value = ("OK", [b""])
        client.logout.return_value = ("BYE", [b""])
        return client

    def test_job_registered(self):
        from nautobot.extras.jobs import get_jobs

        get_jobs(reload=True)
        cls = get_job("redthread_email_import.RedThreadEmailImportJob")
        self.assertIs(cls, RedThreadEmailImportJob)

    @mock.patch.dict(
        os.environ,
        {
            NB_REDTHREAD_IMAP_HOST: "imap.redthread.test",
            NB_REDTHREAD_USERNAME: "ingest@example.com",
            NB_REDTHREAD_PASSWORD: "secret",
        },
        clear=False,
    )
    def test_resolve_credentials_from_env(self):
        u, p = resolve_redthread_imap_credentials()
        self.assertEqual(u, "ingest@example.com")
        self.assertEqual(p, "secret")

    @mock.patch.dict(
        os.environ,
        {
            NB_REDTHREAD_IMAP_HOST: "imap.redthread.test",
            NB_REDTHREAD_USERNAME: "ingest@example.com",
            NB_REDTHREAD_PASSWORD: "secret",
        },
        clear=False,
    )
    @mock.patch.object(redthread_module, "connect_imap_client")
    def test_import_creates_device_from_xlsx(self, mock_connect):
        mime = self._xlsx_mime_bytes()
        mock_connect.return_value = self._mock_imap_client(mime)

        job = RedThreadEmailImportJob()
        job.user = self.user
        job.logger = logging.getLogger("test_redthread_excel_run")
        job.run(
            imap_host="",
            imap_port=993,
            imap_use_ssl=True,
            imap_mailbox="INBOX",
            imap_search_criteria="UNSEEN",
            worksheet_year=0,
            default_location_name=REDTHREAD_DEFAULT_LOCATION_NAME,
            fail_on_missing_device_type=False,
            sender_allowlist="",
            max_messages=25,
            max_attachment_bytes=5_000_000,
            roll_back_if_error=True,
            skip_existing_serial=True,
            mark_seen_on_success=True,
            move_to_mailbox="",
            dryrun=False,
        )

        created = Device.objects.get(serial="RT-IMAP-SN-1")
        self.assertEqual(created.location_id, self.template.location_id)
        self.assertEqual(created.device_type_id, self.template.device_type_id)

    def test_missing_password_raises(self):
        with mock.patch("nautobot.extras.test_jobs.redthread_email_import.os.getenv", return_value=""):
            with self.assertRaises(RunJobTaskFailed):
                resolve_redthread_imap_credentials()
